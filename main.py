import requests, re, html
from docx import Document
from docx.shared import Inches, Pt
from datetime import datetime
from typing import List, Dict, Any

SONAR_URL = "YOUR_SONARQUBE_URL"
PROJECT_KEY = "YOUR_PROJECT_KEY"
TOKEN = "YOUR_TOKEN"

AUTH = (TOKEN, "")
HEADERS = {"Accept": "application/json"}

FILTER_CONFIG = {
    "security": {
        "enabled": True,
        "severities": ["BLOCKER", "HIGH", "MEDIUM"],
        "statuses": ["OPEN"],
        "filename": "sonarqube_security_report.docx"
    },
    "reliability": {
        "enabled": True,
        "severities": ["BLOCKER", "HIGH"],
        "statuses": ["OPEN"],
        "filename": "sonarqube_reliability_report.docx"
    },
    "maintainability": {
        "enabled": True,
        "severities": ["BLOCKER", "HIGH"],
        "statuses": ["OPEN"],
        "filename": "sonarqube_maintainability_report.docx"
    },
    "security_hotspots": {
        "enabled": True,
        "statuses": ["TO_REVIEW"],
        "vulnerability_probabilities": ["HIGH", "MEDIUM"],
        "filename": "sonarqube_security_hotspots_report.docx"
    },
    "title_filters": {
        "skip_prefixes": ["Filename"],
        "skip_contains": [],
        "skip_exact": []
    },
    "source_code": {
        "context_lines": 15
    }
}

SECURITY_CATEGORY_MAP = {
    "dos": "Denial of Service (DoS)",
    "sql-injection": "SQL Injection",
    "xss": "Cross-Site Scripting (XSS)",
    "csrf": "Cross-Site Request Forgery (CSRF)",
    "path-traversal": "Path Traversal",
    "command-injection": "Command Injection",
    "ldap-injection": "LDAP Injection",
    "xpath-injection": "XPath Injection",
    "log-injection": "Log Injection",
    "weak-cryptography": "Weak Cryptography",
    "auth": "Authentication",
    "insecure-conf": "Insecure Configuration",
    "file-manipulation": "File Manipulation",
    "others": "Others"
}

def expand_security_category(category: str) -> str:
    if not category:
        return ""
    category_lower = category.lower().strip()
    if category_lower in SECURITY_CATEGORY_MAP:
        return SECURITY_CATEGORY_MAP[category_lower]
    for key, value in SECURITY_CATEGORY_MAP.items():
        if key in category_lower or category_lower in key:
            return value
    return category

def strip_html(raw: str) -> str:
    if not raw:
        return "N/A"
    clean = re.sub(r"<(br|/p|/div|/h\d)>", "\n", raw, flags=re.I)
    clean = re.sub(r"<[^>]+>", "", clean)
    clean = html.unescape(clean)
    return re.sub(r"\n{2,}", "\n", clean).strip()

def fetch_data(endpoint: str, params: Dict[str, Any], key: str) -> List[Dict[str, Any]]:
    items, page, size = [], 1, 500
    while True:
        params.update({"ps": size, "p": page})
        r = requests.get(f"{SONAR_URL}{endpoint}", auth=AUTH, params=params, headers=HEADERS, verify=False)
        if r.status_code != 200:
            print(f"Request error {endpoint}: {r.text}")
            break
        batch = r.json().get(key, [])
        items.extend(batch)
        if len(batch) < size:
            break
        page += 1
    return items

def fetch_issues_by_quality(quality_type: str) -> List[Dict[str, Any]]:
    config = FILTER_CONFIG[quality_type.lower()]
    return fetch_data("/api/issues/search", {
        "componentKeys": PROJECT_KEY,
        "impactSoftwareQualities": quality_type.upper(),
        "impactSeverities": ",".join(config["severities"]),
        "issueStatuses": ",".join(config["statuses"])
    }, "issues")

def fetch_security_hotspots() -> List[Dict[str, Any]]:
    config = FILTER_CONFIG["security_hotspots"]
    params = {
        "project": PROJECT_KEY,
        "status": ",".join(config["statuses"])
    }
    if "vulnerability_probabilities" in config:
        params["vulnerabilityProbability"] = ",".join(config["vulnerability_probabilities"])
    return fetch_data("/api/hotspots/search", params, "hotspots")

def get_project_name() -> str:
    r = requests.get(f"{SONAR_URL}/api/components/show", auth=AUTH, headers=HEADERS,
                     params={"component": PROJECT_KEY}, verify=False)
    return r.json().get("component", {}).get("name", PROJECT_KEY) if r.status_code == 200 else PROJECT_KEY

def get_analysis_date() -> str:
    r = requests.get(f"{SONAR_URL}/api/project_analyses/search", auth=AUTH, headers=HEADERS,
                     params={"project": PROJECT_KEY}, verify=False)
    if r.status_code == 200:
        analyses = r.json().get("analyses", [])
        if analyses and (d := analyses[0].get("date")):
            return datetime.strptime(d, "%Y-%m-%dT%H:%M:%S%z").strftime("%d %B %Y, %H:%M")
    return ""

def get_source_snippet(component_key: str, line: int, context: int = None) -> str:
    if context is None:
        context = FILTER_CONFIG["source_code"]["context_lines"]
    if not component_key or not line:
        return "N/A"
    try:
        r = requests.get(f"{SONAR_URL}/api/sources/show", auth=AUTH, headers=HEADERS,
                        params={"key": component_key, "from": max(1, line - context), "to": line + context}, verify=False)
        if r.status_code != 200:
            return "N/A"
        data = r.json()
        sources = []
        if isinstance(data, list):
            sources = [item if isinstance(item, dict) and "code" in item else {"code": item[1]} if isinstance(item, list) and len(item) >= 2 else {} for item in data]
        elif isinstance(data, dict):
            if component_key in data and isinstance(data[component_key], dict) and "sources" in data[component_key]:
                sources = data[component_key]["sources"]
            elif "sources" in data:
                sources = data["sources"]
            else:
                sources = next((data[k]["sources"] for k in data if isinstance(data[k], dict) and "sources" in data[k]), [])
        if sources:
            code_lines = [s.get("code", "") if isinstance(s, dict) else s[1] if isinstance(s, list) and len(s) >= 2 else str(s) for s in sources]
            result = "\n".join(code_lines).strip()
            return result if result else "N/A"
    except:
        pass
    return "N/A"

def get_rule_details(rule_key: str, component_key: str = "", line: int = None) -> Dict[str, str]:
    r = requests.get(f"{SONAR_URL}/api/rules/show", auth=AUTH, headers=HEADERS,
                     params={"key": rule_key}, verify=False)
    where = how = category = ""
    if r.status_code == 200:
        rule_data = r.json().get("rule", {})
        desc_sections = rule_data.get("descriptionSections", [])
        
        if "securityStandards" in rule_data:
            for standards_list in rule_data.get("securityStandards", {}).values():
                if isinstance(standards_list, list) and standards_list:
                    category = standards_list[0].get("category", "") or standards_list[0].get("name", "")
                    if category:
                        break
        if not category:
            category = rule_data.get("securityCategory", "")
        
        for s in desc_sections:
            key = s.get("key", "")
            content = s.get("content", "")
            if key == "root_cause":
                where = strip_html(content)
            elif key == "how_to_fix":
                how = strip_html(content)
            elif key == "introduction" and not where:
                where = strip_html(content)
            elif key == "resources" and not how:
                how = strip_html(content)
        
        if not where:
            where = strip_html(rule_data.get("htmlDesc", "") or rule_data.get("mdDesc", ""))
        if not how and not where:
            where = rule_data.get("name", "N/A")
    
    if not where or where == "N/A" or len(where.strip()) == 0:
        where = get_source_snippet(component_key, line)
    
    category = expand_security_category(category)
    
    return {"where": where or "N/A", "how": how or "N/A", "category": category or ""}

def set_font(cell, font="Aptos", size=9):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name, r.font.size = font, Pt(size)

def add_headers(table, headers: List[str]):
    cells = table.rows[0].cells
    for i, text in enumerate(headers):
        cells[i].text = text
        for p in cells[i].paragraphs:
            p.alignment = 1
            if p.runs:
                p.runs[0].bold = True
        set_font(cells[i])

def add_issues(table, issues: List[Dict[str, Any]], widths, quality_type: str = "SECURITY"):
    row_number = 1
    for issue in issues:
        message = issue.get("message", "N/A")
        if should_skip_message(message):
            continue
        
        rule_info = get_rule_details(issue.get("rule", ""), issue.get("component", ""), issue.get("line"))
        comp = issue.get("component", "").split(":")[-1] if isinstance(issue.get("component"), str) else "N/A"
        data = [str(row_number), get_severity(issue, quality_type), comp, message, rule_info["where"], rule_info["how"]]
        add_table_row(table, data, widths, 3)
        row_number += 1

def add_hotspots(table, hotspots: List[Dict[str, Any]], widths):
    row_number = 1
    for hotspot in hotspots:
        message = hotspot.get("message", "N/A")
        if should_skip_message(message):
            continue
        
        rule_info = get_rule_details(hotspot.get("ruleKey", ""), hotspot.get("component", ""), hotspot.get("line"))
        category = expand_security_category(hotspot.get("securityCategory", "") or rule_info.get("category", ""))
        comp = hotspot.get("component", "").split(":")[-1] if isinstance(hotspot.get("component"), str) else "N/A"
        level = hotspot.get("vulnerabilityProbability", hotspot.get("reviewPriority", "N/A"))
        data = [str(row_number), level, category or "N/A", comp, message, rule_info["where"], rule_info["how"]]
        add_table_row(table, data, widths, 4)
        row_number += 1

def create_quality_report(doc, quality_type: str, all_issues: List[Dict[str, Any]], severity_levels: List[str]):
    severity_names = {"BLOCKER": "Blocker", "HIGH": "High", "MEDIUM": "Medium"}
    for severity_level in severity_levels:
        filtered_issues = [issue for issue in all_issues if get_severity(issue, quality_type) == severity_level]
        if filtered_issues:
            doc.add_paragraph(f"\n{severity_names[severity_level]} Severity ({len(filtered_issues)})", style="Heading 2")
            widths = [Inches(0.4), Inches(0.8), Inches(1), Inches(1), Inches(2), Inches(2.5)]
            table, _ = setup_table(doc, ["№", "Level", "Path", "Title", "Where is the issue", "How can I fix that"], widths)
            add_issues(table, filtered_issues, widths, quality_type)

def create_hotspots_report(doc, all_hotspots: List[Dict[str, Any]], probability_levels: List[str]):
    probability_names = {"HIGH": "High", "MEDIUM": "Medium", "LOW": "Low"}
    for prob_level in probability_levels:
        filtered_hotspots = [h for h in all_hotspots if h.get("vulnerabilityProbability", h.get("reviewPriority", "")) == prob_level]
        if filtered_hotspots:
            doc.add_paragraph(f"\n{probability_names.get(prob_level, prob_level)} Probability ({len(filtered_hotspots)})", style="Heading 2")
            widths = [Inches(0.4), Inches(0.8), Inches(0.8), Inches(1), Inches(1), Inches(1.5), Inches(2.5)]
            table, _ = setup_table(doc, ["№", "Level", "Category", "Path", "Title", "Where is the issue", "How can I fix that"], widths)
            add_hotspots(table, filtered_hotspots, widths)

def create_single_report(quality_key: str, quality_name: str, project_name: str, analysis_date: str):
    config = FILTER_CONFIG[quality_key]
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(0.5)
    doc.add_paragraph(f"SonarQube {quality_name} Report", style="Title")
    doc.add_paragraph(f"Project: {project_name}", style="Normal")
    if analysis_date:
        doc.add_paragraph(f"Analysis Date: {analysis_date}", style="Normal")
    
    severities_str = ", ".join([s.capitalize() for s in config["severities"]])
    doc.add_paragraph(f"\nSoftware Quality - {quality_name}\nSeverity - {severities_str}", style="Heading 1")
    
    issues = fetch_issues_by_quality(quality_key)
    create_quality_report(doc, quality_name.upper(), issues, config["severities"])
    
    if not issues:
        doc.add_paragraph(f"No {quality_key} issues found.", style="Normal")
    
    if quality_key == "security" and FILTER_CONFIG["security_hotspots"]["enabled"]:
        hotspots = fetch_security_hotspots()
        if hotspots:
            doc.add_paragraph("\nSecurity Hotspots", style="Heading 1")
            hotspots_config = FILTER_CONFIG["security_hotspots"]
            probabilities_str = ", ".join([p.capitalize() for p in hotspots_config.get("vulnerability_probabilities", [])])
            doc.add_paragraph(f"Vulnerability Probability - {probabilities_str}", style="Normal")
            create_hotspots_report(doc, hotspots, hotspots_config.get("vulnerability_probabilities", ["HIGH", "MEDIUM"]))
        else:
            doc.add_paragraph("\nSecurity Hotspots", style="Heading 1")
            doc.add_paragraph("No security hotspots found.", style="Normal")
    
    doc.save(config["filename"])
    print(f"Report created: {config['filename']}")

def should_skip_message(message: str) -> bool:
    title_filters = FILTER_CONFIG["title_filters"]
    if any(message.startswith(prefix) for prefix in title_filters["skip_prefixes"]):
        return True
    if any(contains in message for contains in title_filters["skip_contains"]):
        return True
    if message in title_filters["skip_exact"]:
        return True
    return False

def add_table_row(table, data, widths, align_threshold=3):
    row = table.add_row().cells
    for i, val in enumerate(data):
        row[i].text = str(val) if val else "N/A"
        row[i].width = widths[i]
        for p in row[i].paragraphs:
            p.alignment = 0 if i > align_threshold else 1
        set_font(row[i])

def get_severity(issue: Dict[str, Any], quality_type: str) -> str:
    for impact in issue.get("impacts", []):
        if impact.get("softwareQuality") == quality_type:
            return impact.get("severity", "N/A")
    return "N/A"

def setup_table(doc, headers: List[str], widths: List) -> tuple:
    table = doc.add_table(rows=1, cols=len(headers), style="Table Grid")
    table.autofit = False
    add_headers(table, headers)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w
    return table, widths

def generate_report():
    project_name = get_project_name()
    analysis_date = get_analysis_date()
    
    reports = [
        ("security", "Security"),
        # ("reliability", "Reliability"),
        # ("maintainability", "Maintainability")
    ]
    
    for quality_key, quality_name in reports:
        if FILTER_CONFIG[quality_key]["enabled"]:
            create_single_report(quality_key, quality_name, project_name, analysis_date)

if __name__ == "__main__":
    generate_report()